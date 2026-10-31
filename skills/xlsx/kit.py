# Copyright (c) 2026 Abdulhamed Samy Eid. All rights reserved.
# SPDX-License-Identifier: LicenseRef-Proprietary
"""XlsxWriter kit: styles, tables, charts, print setup, right-to-left, and finish() (recalc + QA contract + validate).

    sys.path.insert(0, "<skills-root>/xlsx"); from kit import Book

finish() is two-pass when the workbook has charts: XlsxWriter writes chart caches from the values it knows, so
formula cells would chart as zeros. Pass 1 saves, recalculates with LibreOffice and stores the computed values in a
JSON file; finish() then re-runs the generator script (same arguments, env KIT_XLSX_VALUES set); in that second
process every formula is written with value=<computed>, and the final workbook is saved and validated.
"""
from __future__ import annotations

import colorsys
import datetime
import json
import math
import os
import re
import subprocess
import sys
from pathlib import Path

import xlsxwriter
from xlsxwriter.utility import xl_cell_to_rowcol, xl_col_to_name, xl_rowcol_to_cell

HERE = Path(__file__).resolve().parent
INK, MUTED, RULE = "#1F2933", "#6B7785", "#D9E2EC"
INPUT_INK, INPUT_FILL = "#1F4FD8", "#F1F6FD"
LABELS = {False: {"Input": "Blue cells are inputs", "page": "Page &P of &N"},
          True: {"Input": "الخانات الزرقاء للإدخال", "page": "صفحة &P من &N"}}
# "money" shows whole dollars unless the table's typed-in numbers have cents (then "cents"); "money0" and "cents" force one.
NF = {"money": "$#,##0;($#,##0)", "money0": "$#,##0;($#,##0)", "cents": "$#,##0.00;($#,##0.00)", "int": "#,##0",
      "dec": "#,##0.00", "pct": "0.0%", "date": "mmm d, yyyy", "year": "0", "text": None}
MARKERS = ["circle", "diamond", "square", "triangle", "x"]
DASHES = ["solid", "dash", "round_dot", "long_dash", "dash_dot"]
A4 = (595.0, 842.0)  # points


def _rotate(hex_color: str, degrees: float) -> str:
    """The same colour moved round the wheel and quietened: chart series that sit beside the accent without shouting."""
    r, g, b = [int(hex_color[i:i + 2], 16) / 255 for i in (1, 3, 5)]
    h, _, s = colorsys.rgb_to_hls(r, g, b)
    r, g, b = colorsys.hls_to_rgb((h + degrees / 360) % 1, 0.52, min(0.4, max(0.25, s)))
    return "#" + "".join(f"{round(v * 255):02X}" for v in (r, g, b))


ALERT, ALERT_FILL = "#B42318", "#FDE7E4"


def _ramp(hex_color: str) -> dict:
    """One hue from dark to faint: headers and charts stay in one family, and only an alert leaves it."""
    r, g, b = [int(hex_color[i:i + 2], 16) / 255 for i in (1, 3, 5)]
    h, _, s = colorsys.rgb_to_hls(r, g, b)
    tone = lambda light, sat: "#" + "".join(f"{round(v * 255):02X}" for v in colorsys.hls_to_rgb(h, light, sat))
    return {"deep": tone(0.22, min(s, 0.45)), "accent": hex_color, "mid": tone(0.6, s * 0.6), "soft": tone(0.78, s * 0.5), "pale": tone(0.9, 0.4), "faint": tone(0.97, 0.3)}


def _px(width: float, digit: float = 7) -> int:
    """Column width in pixels. Excel and XlsxWriter use 7 px per character; LibreOffice renders about 7.4."""
    return int(width * digit + 0.5) + 5 if width >= 1 else int(width * 12 + 0.5)


def _is_date(spec: str | None) -> bool:
    f = NF.get(spec, spec) if spec else None
    return bool(f) and not any(ch in f for ch in "#0?") and any(ch in f.lower() for ch in "ymd")


def _count(ref: str) -> int:
    """Number of cells in a range reference such as =Model!$B$5:$B$10 (6 if it cannot be read)."""
    m = re.search(r"\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)$", ref)
    if not m:
        return 6
    (r0, c0), (r1, c1) = xl_cell_to_rowcol(m[1] + m[2]), xl_cell_to_rowcol(m[3] + m[4])
    return (r1 - r0 + 1) * (c1 - c0 + 1)


def _quote(name: str) -> str:
    """Sheet-name prefix for formulas. Quote only when needed: LibreOffice rewrites needless quotes and recalc rejects that."""
    return name if name.replace("_", "a").isalnum() and not name[0].isdigit() else "'" + name.replace("'", "''") + "'"


class Table:
    def __init__(self, sheet: "Sheet", r0: int, c0: int, headers: list[str], n: int, total_row: int | None,
                 spec: list | None = None):
        self.sheet, self.r0, self.c0, self.headers, self.n, self.total_row = sheet, r0, c0, headers, n, total_row
        self.spec = spec or [None] * len(headers)

    def _idx(self, col: int | str) -> int:
        return self.headers.index(col) if isinstance(col, str) else col

    def col(self, col: int | str) -> str:
        """Absolute range of one column's data rows (no total), for charts or inside formulas: "=SUM(" + t.col(2) + ")"."""
        j = self._idx(col)
        L = xl_col_to_name(self.c0 + j)
        ref = f"{_quote(self.sheet.name)}!${L}${self.r0 + 2}:${L}${self.r0 + 1 + self.n}"
        if _is_date(self.spec[j]):  # chart category axis reuses the column's date format
            self.sheet.book.date_refs[ref] = "mmm yy" if self.spec[j] == "date" else NF.get(self.spec[j], self.spec[j])
        return ref

    def cell(self, col: int | str, row: int = 0) -> str:
        """Absolute reference of one data cell; row 0 is the first data row, -1 the last."""
        L = xl_col_to_name(self.c0 + self._idx(col))
        return f"{_quote(self.sheet.name)}!${L}${self.r0 + 2 + (row % self.n)}"

    def total(self, col: int | str) -> str:
        """Absolute reference of a column's total cell, e.g. for another sheet's formula."""
        L = xl_col_to_name(self.c0 + self._idx(col))
        return f"{_quote(self.sheet.name)}!${L}${self.total_row + 1}"

    def rule(self, col: int | str, op: str, value, cells: bool = False, **fmt) -> None:
        """Format every data row (or only that column's cells with cells=True) whose `col` value passes a test:
        t.rule("Margin", "<", 0), t.rule("Status", "=", "Late", bg_color="#FFF3CD"). op: < <= > >= = <>. Text is quoted for you.
        Default format: the alert colours; any format options replace them. Rows follow edits, so it needs no helper column."""
        j = self._idx(col)
        L = xl_col_to_name(self.c0 + j)
        v = f'"{value}"' if isinstance(value, str) and not value.startswith("=") else str(value).lstrip("=")
        f = self.sheet.book.wb.add_format(fmt or {"bg_color": ALERT_FILL, "font_color": ALERT})
        first = self.r0 + 1
        c_a, c_z = (self.c0 + j, self.c0 + j) if cells else (self.c0, self.c0 + len(self.headers) - 1)
        self.sheet.ws.conditional_format(first, c_a, first + self.n - 1, c_z, {"type": "formula", "criteria": f"=${L}{first + 1}{op}{v}", "format": f})


class Sheet:
    def __init__(self, book: "Book", name: str, widths: list[float], grid: bool, fit: str | None):
        self.book, self.name, self.grid, self.fit = book, name, grid, fit
        self.ws = book.wb.add_worksheet(name)
        if book.rtl:
            self.ws.right_to_left()
        self.widths = {i: w for i, w in enumerate(widths)}
        for i, w in self.widths.items():
            self.ws.set_column(i, i, w)
        self.row_h: dict[int, float] = {}
        self.max_r = self.max_c = 0
        self.formulas: dict[str, str] = {}  # address -> formula; a cell written twice counts once
        self.charts = 0
        self.freeze: str | None = None
        self.header_row: int | None = None
        self.area: tuple | None = None
        self.roles: set[str] = set()
        self.heading: dict[int, tuple[str, str]] = {}  # row -> (text, role), re-written merged and wrapped if too long

    # ---- writing
    def _ext(self, r: int, c: int) -> None:
        self.max_r, self.max_c = max(self.max_r, r), max(self.max_c, c)

    def _put(self, r: int, c: int, value, role: str | None, nf: str | None, edge: bool = False, **kw) -> None:
        if role and ":" in role:
            role, nf = role.split(":", 1)
        is_formula = isinstance(value, str) and value.startswith("=")
        addr = xl_rowcol_to_cell(r, c)
        is_date = isinstance(value, datetime.date)
        if role is None:
            if is_formula:
                role = "link" if "!" in value else "calc"
            else:
                role = "in" if is_date or isinstance(value, (int, float)) and not isinstance(value, bool) else "label"
        if is_date and nf is None:
            nf = "date"
        if nf == "money" and isinstance(value, (int, float)) and value != int(value):
            nf = "cents"
        self.roles.add(role)
        fmt = self.book.fmt(role, nf, edge, **kw)
        if is_formula:
            self.ws.write_formula(r, c, value, fmt, self.book.values.get(f"{self.name}!{addr}", 0))
            self.formulas[addr] = value
        else:
            self.ws.write(r, c, value, fmt)
            self.formulas.pop(addr, None)
        self._ext(r, c)

    def put(self, at: str, value, role: str | None = None, nf: str | None = None, **kw) -> None:
        """Write one cell. role: in|calc|link|total|head|label|note ('in:money' sets format too).
        Default role: formula -> calc (link if it has '!'), number -> in, text -> label."""
        r, c = xl_cell_to_rowcol(at)
        self._put(r, c, value, role, nf, **kw)

    def title(self, text: str, sub: str | None = None) -> None:
        self._put(0, 0, text, "title", None)
        self.row_h[0] = 26
        self.heading[0] = (text, "title")
        if sub:
            self._put(1, 0, sub, "sub", None)
            self.row_h[1] = 17
            self.heading[1] = (sub, "sub")

    def legend(self, at: str | None = None) -> None:
        """One line saying what the blue cells are; call it after the tables. Default place: column A, two rows below the
        last written row. Writes nothing if the sheet has no inputs."""
        r, c = xl_cell_to_rowcol(at) if at else (self.max_r + 2, 0)
        if "in" in self.roles:
            self._put(r, c, LABELS[self.book.rtl]["Input"], "note", None, font_color=INPUT_INK)

    def inputs(self, at: str, items: list[tuple]) -> dict[str, str]:
        """Two columns, no header: items = [(label, value, fmt[, note])] -> label in `at`, blue input value one column right
        (numbers, or datetime.date written as a real date), optional muted note after it. Returns {label: "Sheet!$B$4"}."""
        r0, c0 = xl_cell_to_rowcol(at)
        refs = {}
        for i, (label, value, nf, *note) in enumerate(items):
            self._put(r0 + i, c0, label, "label", None, True)
            self._put(r0 + i, c0 + 1, value, "in", nf, True)
            if note:
                self._put(r0 + i, c0 + 2, note[0], "note", None, indent=1)
            self.row_h[r0 + i] = 20
            refs[label] = f"{_quote(self.name)}!${xl_col_to_name(c0 + 1)}${r0 + i + 1}"
        return refs

    def kpi(self, at: str, items: list[tuple]) -> None:
        """Headline tiles across columns: items = [(label, formula, fmt)]."""
        r, c = xl_cell_to_rowcol(at)
        self.row_h[r], self.row_h[r + 1] = 20, 36
        for j, (label, formula, nf) in enumerate(items):
            self._put(r, c + j, label, "kpi_label", None)
            self._put(r + 1, c + j, formula, "kpi", nf)

    def link(self, at: str, sheet: str, text: str | None = None) -> None:
        r, c = xl_cell_to_rowcol(at)
        self.ws.write_url(r, c, f"internal:{_quote(sheet)}!A1", self.book.fmt("hyperlink", None), string=text or sheet)
        self._ext(r, c)

    def table(self, at: str, headers: list[str], rows: list[list], fmts: list[str | None] | None = None,
              total: str | None = None, sums: list[int] | int = (), freeze: bool = True, filt: bool = False,
              rh: float = 20, inputs: bool = True, totals: dict | None = None) -> Table:
        """Header + data rows (+ total row). A string cell starting '=' is a formula; '{r}' is its row number, '{p}' the row above
        (not allowed in the first data row: pass a (first_row_cell, later_rows_cell) tuple instead).
        fmts[j]: money|money0|cents|int|dec|pct|date|year|text, or any Excel format; 'role:fmt' forces a role.
        Plain numbers and dates are styled as inputs (blue); inputs=False styles them as plain data (raw/source tables).
        totals={col: formula} sets a total-row cell by formula (a ratio, a weighted mean): {r} is the total row, {a} the first and
        {z} the last data row, e.g. {3: "=D{r}/C{r}", 4: "=SUMPRODUCT(B{a}:B{z},E{a}:E{z})/SUM(B{a}:B{z})"}."""
        r0, c0 = xl_cell_to_rowcol(at)
        n = len(headers)
        fmts = list(fmts or [None] * n)
        sums = [sums] if isinstance(sums, int) else list(sums)
        totals = {(headers.index(k) if isinstance(k, str) else k): v for k, v in (totals or {}).items()}
        spec = [(f or "").split(":", 1)[-1] or None for f in fmts]
        if any(isinstance(v, (int, float)) and not isinstance(v, bool) and v != int(v) for row in rows for v in row):
            spec = ["cents" if s == "money" else s for s in spec]  # typed-in cents anywhere: all "money" columns show cents
        left = [s in (None, "text") for s in spec]
        for j, h in enumerate(headers):
            # a filter arrow sits at the cell's right edge: right-aligned text goes under it, so filtered headers align left with room for it
            self._put(r0, c0 + j, h, "head", None, **({"align": "left", "indent": 1} if left[j] or filt else {"align": "right"}))
            if filt and self.widths.get(c0 + j, 8.43) < len(str(h)) * 1.15 + 5:
                self.widths[c0 + j] = len(str(h)) * 1.15 + 5
                self.ws.set_column(c0 + j, c0 + j, self.widths[c0 + j])
        self.row_h[r0] = 24
        for i, row in enumerate(rows):
            r = r0 + 1 + i
            self.row_h[r] = rh
            for j, v in enumerate(row):
                if isinstance(v, tuple):
                    v = v[min(i, 1)]
                if isinstance(v, str) and v.startswith("="):
                    if i == 0 and "{p}" in v:
                        raise ValueError(f"{headers[j]!r}: '{{p}}' in the first data row points at the header; "
                                         "use a (first_row_cell, later_rows_cell) tuple")
                    v = v.replace("{r}", str(r + 1)).replace("{p}", str(r))
                f = fmts[j] or ""
                role = f.split(":", 1)[0] if ":" in f else (None if inputs else "label")
                self._put(r, c0 + j, v, role, spec[j], True, **({"indent": 1} if left[j] else {}))
        last = r0 + len(rows)
        total_row = None
        if total is not None:
            total_row = last + 1
            self.row_h[total_row] = rh + 4
            for j in range(n):
                L = xl_col_to_name(c0 + j)
                if j == 0:
                    self._put(total_row, c0, total, "total", None, indent=1)
                elif j in sums:
                    self._put(total_row, c0 + j, f"=SUM({L}{r0 + 2}:{L}{last + 1})", "total", spec[j])
                elif j in totals:
                    f = totals[j].replace("{r}", str(total_row + 1)).replace("{a}", str(r0 + 2)).replace("{z}", str(last + 1))
                    self._put(total_row, c0 + j, f, "total", spec[j])
                else:
                    self._put(total_row, c0 + j, "", "total", None)
        if freeze and self.freeze is None:
            self.ws.freeze_panes(r0 + 1, 0)
            self.freeze = xl_rowcol_to_cell(r0 + 1, 0)
        if self.header_row is None:
            self.header_row = r0
        if filt:
            self.ws.autofilter(r0, c0, last, c0 + n - 1)
        return Table(self, r0, c0, headers, len(rows), total_row, spec)

    def chart(self, chart, rng: str, rows: int = 16) -> None:
        """Place a chart so it exactly fills the cell block, e.g. 'A11:D24' (keeps it inside the print area).
        Columns only ('A:D') puts it `rows` tall, one blank row below everything written so far (charts stack)."""
        a, b = rng.split(":")
        if a.isalpha():
            a, b = f"{a}{self.max_r + 3}", f"{b}{self.max_r + 2 + rows}"
        (r0, c0), (r1, c1) = xl_cell_to_rowcol(a), xl_cell_to_rowcol(b)
        w = sum(_px(self.widths.get(c, 8.43)) for c in range(c0, c1 + 1))
        h = sum(round(self.row_h.get(r, 15) / 0.75) for r in range(r0, r1 + 1))
        chart.set_size({"width": w - 6, "height": h - 6})
        self.ws.insert_chart(r0, c0, chart, {"x_offset": 3, "y_offset": 3, "object_position": 2})
        self._ext(r1, c1)
        self.charts += 1

    # ---- print
    def print_area(self, rng: str) -> None:
        """Override the print area (default: every written cell and chart), e.g. 'A1:D26'."""
        a, b = rng.split(":")
        (r0, c0), (r1, c1) = xl_cell_to_rowcol(a), xl_cell_to_rowcol(b)
        self.area = (r0, c0, r1, c1)

    def _print(self) -> None:
        ws = self.ws
        ws.set_paper(9)
        ws.hide_gridlines(1 if self.grid else 2)  # 1 = screen only; never 0, which prints gridlines
        r0, c0, r1, c1 = self.area or (0, 0, self.max_r, self.max_c)
        self.max_r, self.max_c = r1, c1
        avail = sum(_px(self.widths.get(c, 8.43), 7.5) for c in range(c1 + 1))
        for r, (text, role) in self.heading.items():  # a title or subtitle wider than the print area wraps over it
            per_line = 12.5 if role == "title" else 6.3
            lines = math.ceil(len(text) * per_line / avail)
            if lines > 1 and c1 > 0:
                ws.merge_range(r, 0, r, c1, text, self.book.fmt(role, None, False, text_wrap=True))
                self.row_h[r] = 26 * lines + 6 if role == "title" else 15 * lines + 3
        for r, h in self.row_h.items():
            ws.set_row(r, h)
        ws.print_area(r0, c0, r1, c1)
        ws.center_horizontally()
        ws.set_margins(0.5, 0.5, 0.7, 0.7)
        w = avail * 0.75
        h = sum(self.row_h.get(r, 15) for r in range(self.max_r + 1))
        pw, ph = A4[0] - 72, A4[1] - 100
        # "page": the whole sheet on one page, enlarged up to 160% to fill the width, printed from the top.
        # "width": one page wide, as tall as needed. Auto: page when that needs no more than 40% shrinking.
        fit = self.fit or ("page" if max(min(pw / w, ph / h), min(ph / w, pw / h)) >= 0.6 else "width")
        page = fit == "page"
        cap = 1.6 if page else 1.15
        scale = {}
        for orient, (aw, ah) in {"portrait": (pw, ph), "landscape": (ph, pw)}.items():
            scale[orient] = min(cap, aw / w, ah / h) if page else min(cap, aw / w)
        orient = "portrait" if scale["portrait"] >= scale["landscape"] * (0.87 if page else 0.75) else "landscape"
        ws.set_landscape() if orient == "landscape" else ws.set_portrait()
        ws.set_print_scale(max(20, int(min(cap, scale[orient]) * 97)))
        if not page and self.header_row is not None:
            ws.repeat_rows(self.header_row)
        text, page_no = self.book.footer or self.name, LABELS[self.book.rtl]["page"]
        ws.set_footer(f"&L&9{page_no}&R&9{text}" if self.book.rtl else f"&L&9{text}&R&9{page_no}")
        self.orient = orient


class Dash(Sheet):
    """A dashboard page: equal grid columns on a pale ground, no gridlines. Components are placed by grid column and row, from 1,
    and each is white on the ground: header, nav, select, card, panel, chart, matrix, text. With a select, everything fed by it updates."""

    def __init__(self, book: "Book", name: str, cols: int, col_w: float, row_h: float):
        super().__init__(book, name, [1.5] + [col_w] * cols + [1.5], False, "page")
        self.cols, self.std_h, self.selected = cols, row_h, False
        ground = book.wb.add_format({"bg_color": book.ramp["faint"]})
        self.ws.set_column(0, 0, 1.5, ground)
        self.ws.set_column(1, cols, col_w, ground)
        self.ws.set_column(cols + 1, cols + 1, 1.5, ground)
        self.ws.set_column(cols + 2, 16383, None, None, {"hidden": True})  # unused rows and columns are hidden, so the page reads as one surface
        self.ws.set_default_row(row_h, True)

    def _print(self) -> None:
        for r in range(self.max_r + 2):  # every row in use must be set, or it stays hidden with the rest
            self.row_h.setdefault(r, self.std_h)
        super()._print()

    def _f(self, nf=None, **kw):
        return self.book.fmt("label", nf, False, **kw)

    def _gut(self, first=True, last=True):
        """A thick ground-coloured edge, the gap that keeps neighbouring blocks apart without spending a grid column on it."""
        g = self.book.ramp["faint"]
        return {**({"left": 5, "left_color": g} if first else {}), **({"right": 5, "right_color": g} if last else {})}

    def _block(self, row, col, rows, span, **kw):
        for r in range(row - 1, row - 1 + rows):
            for c in range(col, col + span):
                self.ws.write_blank(r, c, None, self._f(**kw, **self._gut(c == col, c == col + span - 1)))
        self._ext(row - 2 + rows, col + span - 1)

    def _text(self, row, col, span, value, nf=None, rows=1, **kw):
        """A merged run of cells holding text or a live formula."""
        kw.setdefault("align", "left")  # "left" is the start of the line; a right-to-left book flips it
        fmt = self._f(nf, **kw)
        r, c = row - 1, col
        if span > 1 or rows > 1:
            self.ws.merge_range(r, c, r + rows - 1, c + span - 1, "", fmt)
        self._put(r, c, value, "label", nf, **kw)
        self._ext(r + rows - 1, c + span - 1)

    def text(self, col, row, span, value, rows=1, **kw):
        """A note on the page's ground, wrapped; the value may be a formula. kw are format options."""
        kw = {"bg_color": self.book.ramp["faint"], "font_color": MUTED, "font_size": 9, "text_wrap": True, "valign": "top", **kw}
        self._text(row, col, span, value, rows=rows, **kw)

    def header(self, title, sub=None):
        """A band across the top, rows 1 and 2."""
        deep = self.book.ramp["deep"]
        self._text(1, 1, self.cols, title, bg_color=deep, font_color="#FFFFFF", bold=True, font_size=16, indent=1)
        self._text(2, 1, self.cols, sub or "", bg_color=deep, font_color=self.book.ramp["soft"], indent=1, valign="top")
        self.row_h[0], self.row_h[1] = 30, 20

    def nav(self, row, sheets):
        """A strip of links to other sheets across the page. Longer names get the spare columns."""
        ramp, n = self.book.ramp, len(sheets)
        spans = [self.cols // n] * n
        for i in sorted(range(n), key=lambda i: -len(sheets[i]))[:self.cols % n]:
            spans[i] += 1
        fmt, col = self._f(bg_color=ramp["pale"], font_color=ramp["deep"], font_size=9, align="center"), 1
        for name, span in zip(sheets, spans):
            if span > 1:
                self.ws.merge_range(row - 1, col, row - 1, col + span - 1, "", fmt)
            self.ws.write_url(row - 1, col, f"internal:{_quote(name)}!A1", fmt, string=name)
            col += span
        self._ext(row - 1, col - 1)
        self.row_h[row - 1] = 20

    def select(self, col, row, span, label, options, default):
        """A dropdown the reader changes: label above, value below. Returns its absolute reference, for formulas."""
        ramp = self.book.ramp
        self._text(row, col, span, label + "  ▾", bg_color=ramp["faint"], font_color=ramp["deep"], font_size=9, bold=True)
        self._text(row + 1, col, span, default, bg_color="#FFFFFF", font_color=INPUT_INK, bold=True, font_size=12, border=1, border_color=ramp["accent"], align="center")
        self.row_h[row] = 24
        self.ws.data_validation(row, col, row, col, {"validate": "list", "source": options})
        if not self.selected:  # the book opens on the first dropdown, so its arrow is already showing
            self.ws.set_selection(row, col, row, col)
            self.selected = True
        return f"{_quote(self.name)}!${xl_col_to_name(col)}${row + 1}"

    def card(self, col, row, span, label, value, nf=None, note=None, alert=None):
        """One figure that matters: label, value, optional note. 4 rows tall. alert=">1" turns the value red when true."""
        ramp, gut = self.book.ramp, self._gut()
        self._text(row, col, span, label, bg_color="#FFFFFF", font_color=ramp["accent"], font_size=9, bold=True, top=5, top_color=ramp["accent"], indent=1, **gut)
        self._text(row + 1, col, span, value, nf, rows=2, bg_color="#FFFFFF", font_color=ramp["deep"], bold=True, font_size=20, indent=1, align="left", **gut)
        self._text(row + 3, col, span, note or "", bg_color="#FFFFFF", font_color=MUTED, font_size=9, indent=1, **gut)
        if alert:
            op, val = re.fullmatch(r"([<>=!]+)\s*(.+)", alert).groups()
            self.ws.conditional_format(row, col, row + 1, col, {"type": "cell", "criteria": op, "value": val, "format": self.book.wb.add_format({"font_color": ALERT})})

    def panel(self, col, row, span, rows, title):
        """A white block with a title; the title may be a live formula."""
        self._block(row, col, rows, span, bg_color="#FFFFFF")
        self._text(row, col, span, title, bg_color="#FFFFFF", font_color=self.book.ramp["deep"], bold=True, font_size=11, indent=1, bottom=1, bottom_color=self.book.ramp["pale"], **self._gut())
        self.row_h[row - 1] = 24

    def chart(self, chart, col, row, span, rows, title):
        """A panel filled by a chart. Build the chart with b.chart(...); its title comes from the panel."""
        self.panel(col, row, span, rows, title)
        chart.set_title({"none": True})
        chart.set_chartarea({"border": {"none": True}, "fill": {"color": "#FFFFFF"}})
        font = chart.x_axis.get("num_font")
        if font and getattr(chart, "short_axis", False):  # a short axis reads flat; left to itself Excel tilts labels that nearly fit
            font["size"], font["rotation"], chart.x_axis["interval_unit"] = 900, 0, 1  # XlsxWriter holds sizes in hundredths
        w = sum(_px(self.widths.get(c, 8.43)) for c in range(col, col + span))
        h = round(self.std_h / 0.75) * (rows - 1)
        chart.set_size({"width": w - 8, "height": h - 12})
        self.ws.insert_chart(row, col, chart, {"x_offset": 4, "y_offset": 3, "object_position": 2})
        self.charts += 1

    def matrix(self, col, row, title, rows, cols, cell, nf="pct", over=None, label_span=2):
        """A heat table in a panel, one grid column per cell after the row labels: rows and cols are labels, cell(row_index, col_index)
        gives a value or formula. The fill deepens with the value; over=1 marks values above it in red."""
        ramp = self.book.ramp
        start, end = col + label_span, col + label_span + len(cols) - 1
        self.panel(col, row, label_span + len(cols), len(rows) + 3, title)
        self._text(row + 1, col, label_span, "", bg_color="#FFFFFF", bottom=1, bottom_color=ramp["pale"], **self._gut(True, False))
        for j, h in enumerate(cols):
            self._text(row + 1, start + j, 1, h, bg_color="#FFFFFF", font_color=MUTED, bold=True, font_size=9, align="center", bottom=1, bottom_color=ramp["pale"], **self._gut(False, start + j == end))
        for i, name in enumerate(rows):
            r = row + 2 + i
            self._text(r, col, label_span, name, None, bg_color="#FFFFFF", bold=True, indent=1, **self._gut(True, False))
            for j in range(len(cols)):
                self._text(r, start + j, 1, cell(i, j), nf, bg_color="#FFFFFF", align="center", **self._gut(False, start + j == end))
        first, last = row + 1, row + len(rows)
        if over is not None:
            self.ws.conditional_format(first, start, last, end, {"type": "cell", "criteria": ">", "value": over, "stop_if_true": True,
                                       "format": self.book.wb.add_format({"bg_color": ALERT_FILL, "font_color": ALERT, "bold": True})})
        self.ws.conditional_format(first, start, last, end, {"type": "2_color_scale", "min_color": "#FFFFFF", "max_color": ramp["soft"]})


class Book:
    def __init__(self, path: str, accent: str | None = None, footer: str | None = None, font: str = "Arial", rtl: bool = False):
        self.rtl = bool(rtl)
        if not accent:
            raise ValueError('choose the colour of this workbook: Book(path, accent="#2F5D50")')
        self.ramp = _ramp(accent)
        self.out = Path(path).resolve()
        self.out.parent.mkdir(parents=True, exist_ok=True)
        self.raw = self.out.with_name(f".{self.out.stem}.raw.xlsx")
        # XlsxWriter's default is Calibri, which renderers without it substitute; set a font that exists everywhere
        self.font = font
        self.wb = xlsxwriter.Workbook(str(self.raw), {"default_format_properties": {"font_name": font, "font_size": 10}})
        self.accent, self.footer = accent, footer
        self.palette = [self.ramp["accent"], self.ramp["deep"], self.ramp["mid"], self.ramp["soft"], *(_rotate(accent, d) for d in (150, 40, -50, 200))]
        self.sheets: list[Sheet] = []
        self._fmts: dict = {}
        self.wb.set_properties({"title": self.out.stem})
        self.date_refs: dict[str, str] = {}
        values = os.environ.get("KIT_XLSX_VALUES")  # set only in finish()'s second pass
        self.values: dict = json.loads(Path(values).read_text(encoding="utf-8")) if values else {}
        self.pass2 = bool(values)

    def dash(self, name: str, cols: int = 14, col_w: float = 11.5, row_h: float = 18) -> Dash:
        """A dashboard page; see class Dash."""
        s = Dash(self, name, cols, col_w, row_h)
        self.sheets.append(s)
        return s

    def sheet(self, name: str, widths: list[float] = (), grid: bool = True, fit: str | None = None) -> Sheet:
        """grid=False hides the screen gridlines (a chart or summary page); they never print. fit='page'|'width' overrides the print mode."""
        s = Sheet(self, name, list(widths), grid, fit)
        self.sheets.append(s)
        return s

    def fmt(self, role: str, nf: str | None = None, edge: bool = False, **kw):
        key = (role, nf, edge, tuple(sorted(kw.items())))
        if key in self._fmts:
            return self._fmts[key]
        a = self.accent
        d = {"font_name": self.font, "font_color": INK, "valign": "vcenter"}
        d.update({
            "label": {}, "calc": {}, "link": {},
            "in": {"font_color": INPUT_INK, "bg_color": INPUT_FILL},
            "head": {"bold": True, "font_color": self.ramp["deep"], "bg_color": self.ramp["pale"], "bottom": 2, "bottom_color": a},
            "total": {"bold": True, "top": 1, "top_color": INK},
            "note": {"font_color": MUTED, "font_size": 9},
            "title": {"bold": True, "font_size": 15},
            "sub": {"font_color": MUTED},
            "kpi_label": {"font_color": MUTED, "font_size": 9, "indent": 1},
            "kpi": {"bold": True, "font_size": 16, "align": "left", "indent": 1},
            "hyperlink": {"font_color": INPUT_INK, "underline": 1},
        }[role])
        if nf and NF.get(nf, nf):
            d["num_format"] = NF.get(nf, nf)
        if edge:  # gridlines do not print, so data rows carry a pale rule of their own for paper
            d.update(bottom=1, bottom_color=RULE)
        d.update(kw)
        if self.rtl:  # the sheet is mirrored: text starts at the right, and numbers stay right-aligned as Excel sets them
            d["reading_order"] = 2
            if d.get("align") == "left":
                d["align"] = "right"
        self._fmts[key] = self.wb.add_format(d)
        return self._fmts[key]

    def chart(self, kind: str, title: str, cats: str, series: list[tuple], nf: str = "money0",
              labels: bool | None = None, cat_nf: str | None = None, y_min: float | None = None, legend: bool | None = None):
        """kind: column|bar|line|area|pie|doughnut, optionally 'column:stacked'. series = [(name, values_ref[, opts])];
        cats/values come from Table.col(). Series opts {'as': 'line'} draws that series as a line on a secondary axis
        (as a percentage) over the columns. Pie/doughnut: one series, percentage labels, one colour per slice. labels: value
        labels, default on for a single column/bar series only. cat_nf: category-axis number format, default taken from a
        date-formatted table column ('date' shows 'mmm yy'; a custom date format is used as is)."""
        typ, _, sub = kind.partition(":")
        chart = self.wb.add_chart({"type": typ, "subtype": sub} if sub else {"type": typ})
        overlay = self.wb.add_chart({"type": "line"}) if any(len(s) > 2 and s[2].get("as") == "line" for s in series) else None
        many = len(series) > 1
        nf_str = NF.get(nf, nf)
        pie = typ in ("pie", "doughnut")
        for i, item in enumerate(series):
            name, values = item[:2]
            opts = dict(item[2]) if len(item) > 2 else {}
            as_line = opts.pop("as", None) == "line"
            color = self.palette[i % len(self.palette)]
            d = {"name": name, "categories": cats, "values": values}
            if pie:
                d["points"] = [{"fill": {"color": self.palette[k % len(self.palette)]}} for k in range(_count(values))]
            elif typ == "line" or as_line:
                d["line"] = {"color": color, "width": 2.5, "dash_type": DASHES[i % 5]}
                d["marker"] = {"type": MARKERS[i % 5], "size": 7, "fill": {"color": color}, "border": {"color": color}}
            else:
                d["fill"], d["border"] = {"color": color}, {"none": True}
                if typ in ("column", "bar"):
                    d["gap"] = 70
            if pie:
                d["data_labels"] = {"percentage": True, "position": "outside_end", "font": {"size": 10, "color": INK}}
            elif labels or (labels is None and not many and typ in ("column", "bar")):
                d["data_labels"] = {"value": True, "num_format": nf_str, "font": {"size": 10, "color": INK}}
            if as_line:
                d["y2_axis"] = True
            d.update(opts)
            (overlay if as_line else chart).add_series(d)
        chart.short_axis = typ != "bar" and not pie and _count(cats) <= 14
        chart.set_title({"name": title, "name_font": {"size": 12, "bold": True, "color": INK}})
        chart.set_chartarea({"border": {"none": True}})
        if not pie:
            cat = {"num_font": {"size": 10, "color": INK}, "line": {"color": RULE}, "label_position": "low"}  # below negative bars, not across them
            if cat_nf or cats.lstrip("=") in self.date_refs:
                cat["num_format"] = cat_nf or self.date_refs[cats.lstrip("=")]
            val = {"num_format": nf_str, "num_font": {"size": 9, "color": MUTED}, "line": {"none": True},
                   "major_gridlines": {"visible": True, "line": {"color": RULE, "width": 0.75}}}
            if y_min is not None:
                val["min"] = y_min
            if self.rtl:  # first category at the right, the value axis on the right, bars growing from the right
                if typ == "bar":
                    val["reverse"] = True
                else:
                    cat["reverse"] = True
            if typ == "bar":  # XlsxWriter's x axis is the horizontal one on bar charts
                chart.set_x_axis(val)
                chart.set_y_axis({**cat, "reverse": True})  # first row at the top
            else:
                chart.set_x_axis(cat)
                chart.set_y_axis(val)
        if overlay:
            overlay.set_y2_axis({"num_format": NF["pct"], "num_font": {"size": 9, "color": MUTED}, "line": {"none": True}})
            chart.combine(overlay)
        chart.set_legend({"position": "bottom", "font": {"size": 10}} if (many or pie if legend is None else legend) else {"none": True})
        return chart

    def _recalc(self, dst: Path) -> None:
        run = subprocess.run([sys.executable, str(HERE / "scripts" / "recalculate_workbook.py"), str(self.raw), str(dst)],
                             capture_output=True, text=True, timeout=240)
        if run.returncode != 0:
            raise SystemExit(f"recalculation failed: {(run.stdout + run.stderr)[:1500]}")

    def _rerun_with_values(self, calculated: Path) -> None:
        """Pass 1 -> 2: store the computed value of every formula cell, re-run this script, and exit with its status."""
        from openpyxl import load_workbook
        from openpyxl.utils.datetime import to_excel

        book = load_workbook(calculated, data_only=True)
        values = {}
        for s in self.sheets:
            for addr in s.formulas:
                v = book[s.name][addr].value
                if v is not None:
                    values[f"{s.name}!{addr}"] = to_excel(v) if hasattr(v, "isoformat") else v
        book.close()
        store = self.out.with_name(f".{self.out.stem}.values.json")
        store.write_text(json.dumps(values), encoding="utf-8")
        for p in (self.raw, calculated):
            p.unlink(missing_ok=True)
        try:
            done = subprocess.run([sys.executable, *sys.argv], env={**os.environ, "KIT_XLSX_VALUES": str(store)}, timeout=600)
        finally:
            store.unlink(missing_ok=True)
        raise SystemExit(done.returncode)

    def finish(self, expect: dict | None = None, extra: dict | None = None) -> None:
        """Close, recalculate with LibreOffice, write <out>.qa.json, run the validator. Raises on any failure.
        With charts and formulas it re-runs the generator once (see module doc), so the script must be re-runnable.
        expect = {"Sheet!B9": 2400} asserts supplied or independently known results (t.total(3) works as a key).
        extra = more QA contract keys (requiredTables, keyRelations, borderRanges ...), merged over the generated ones."""
        for s in self.sheets:
            s._print()
        self.wb.close()
        if self.pass2:
            os.replace(self.raw, self.out)
        elif any(s.charts for s in self.sheets) and any(s.formulas for s in self.sheets) and Path(sys.argv[0]).is_file():
            pass1 = self.out.with_name(f".{self.out.stem}.pass1.xlsx")
            self._recalc(pass1)
            self._rerun_with_values(pass1)
        else:
            self._recalc(self.out)
            self.raw.unlink(missing_ok=True)
        formulas = [(s.name, a, f) for s in self.sheets for a, f in s.formulas.items()]
        contract = {
            "schemaVersion": 1,
            "exactSheets": [s.name for s in self.sheets],
            "minimumFormulaCells": len(formulas),
            "requireFormulaCaches": True,
            "requiredFormulas": [{"cell": f"{n}!{a}", "formula": f} for n, a, f in formulas[:40]],
            "expectedValues": [{"cell": k.replace("$", ""), "value": v, "tolerance": 0.01} for k, v in (expect or {}).items()],
            "requiredCharts": [{"sheet": s.name, "minimum": s.charts} for s in self.sheets if s.charts],
            "freezePanes": [{"sheet": s.name, "cell": s.freeze} for s in self.sheets if s.freeze],
            "printSettings": [{"sheet": s.name, "orientation": s.orient, "requirePrintArea": True} for s in self.sheets],
            "sheetViews": [{"sheet": s.name, "showGridLines": s.grid} for s in self.sheets],
        }
        contract.update(extra or {})
        qa = Path(str(self.out) + ".qa.json")
        qa.write_text(json.dumps(contract, indent=2), encoding="utf-8")
        check = subprocess.run([sys.executable, str(HERE / "scripts" / "validate_workbook.py"), str(self.out), "--json"],
                               capture_output=True, text=True, timeout=240)
        report = json.loads(check.stdout)
        if not report["ok"]:
            raise SystemExit("validation failed:\n" + "\n".join(f"  {i.get('code')} {i.get('location', '')} {i.get('message')}"
                                                                for i in report["issues"][:20]))
        print(f"OK {self.out} formulas={len(formulas)} charts={sum(s.charts for s in self.sheets)} validator=ok")
