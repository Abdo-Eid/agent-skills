#!/usr/bin/env python3
# SPDX-License-Identifier: MIT
"""Recalculate an XLSX through isolated LibreOffice and verify formula caches."""

from __future__ import annotations

import argparse
import hashlib
import json
import posixpath
import shutil
import subprocess
import sys
import tempfile
import zipfile
from copy import copy
from pathlib import Path
from typing import Any
from xml.etree import ElementTree


PROFILE_XML = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry">
  <item oor:path="/org.openoffice.Office.Calc/Formula/Load">
    <prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop>
  </item>
</oor:items>
"""

SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
DOCUMENT_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


class RuntimeUnavailableError(RuntimeError):
    pass


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def check_zip(path: Path) -> None:
    if not path.is_file() or path.stat().st_size <= 0:
        raise RuntimeError("LibreOffice did not produce a non-empty output workbook")
    if not zipfile.is_zipfile(path):
        raise RuntimeError("LibreOffice output is not a valid XLSX ZIP package")
    with zipfile.ZipFile(path) as archive:
        broken = archive.testzip()
        if broken:
            raise RuntimeError(f"LibreOffice output contains a corrupt ZIP entry: {broken}")
        required = {"[Content_Types].xml", "xl/workbook.xml"}
        missing = sorted(required.difference(archive.namelist()))
        if missing:
            raise RuntimeError(f"LibreOffice output is missing required part: {missing[0]}")


def workbook_report(path: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeUnavailableError("openpyxl is missing from the baked runtime") from error

    formula_book = load_workbook(path, read_only=True, data_only=False, keep_links=True)
    cached_book = load_workbook(path, read_only=True, data_only=True, keep_links=True)
    formulas: dict[str, str] = {}
    missing: list[str] = []
    error_values: list[str] = []
    try:
        for formula_sheet, cached_sheet in zip(
            formula_book.worksheets, cached_book.worksheets
        ):
            for formula_row, cached_row in zip(
                formula_sheet.iter_rows(), cached_sheet.iter_rows()
            ):
                for formula_cell, cached_cell in zip(formula_row, cached_row):
                    if formula_cell.data_type != "f":
                        continue
                    location = f"{formula_sheet.title}!{formula_cell.coordinate}"
                    formulas[location] = str(formula_cell.value)
                    if cached_cell.value is None:
                        missing.append(location)
                    elif cached_cell.data_type == "e" or (
                        isinstance(cached_cell.value, str)
                        and cached_cell.value.startswith("#")
                    ):
                        error_values.append(f"{location}={cached_cell.value}")
    finally:
        formula_book.close()
        cached_book.close()
    return {
        "formulas": formulas,
        "formulaCells": len(formulas),
        "cachedValues": len(formulas) - len(missing),
        "missingCachedValues": missing,
        "cachedErrorValues": error_values,
    }


def worksheet_parts(archive: zipfile.ZipFile) -> dict[str, str]:
    workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    relationships = ElementTree.fromstring(
        archive.read("xl/_rels/workbook.xml.rels")
    )
    targets = {
        relationship.attrib["Id"]: relationship.attrib["Target"]
        for relationship in relationships.findall(f"{{{PACKAGE_REL_NS}}}Relationship")
    }
    parts: dict[str, str] = {}
    for sheet in workbook.findall(f".//{{{SHEET_NS}}}sheet"):
        relationship_id = sheet.attrib[f"{{{DOCUMENT_REL_NS}}}id"]
        target = targets[relationship_id]
        part = target.lstrip("/") if target.startswith("/") else posixpath.normpath(
            posixpath.join("xl", target)
        )
        parts[sheet.attrib["name"]] = part
    return parts


def transplant_formula_caches(source: Path, calculated: Path, output: Path) -> int:
    ElementTree.register_namespace("", SHEET_NS)
    with zipfile.ZipFile(source) as source_zip, zipfile.ZipFile(calculated) as calculated_zip:
        source_parts = worksheet_parts(source_zip)
        calculated_parts = worksheet_parts(calculated_zip)
        if source_parts.keys() != calculated_parts.keys():
            raise RuntimeError("LibreOffice changed worksheet names during recalculation")

        replacements: dict[str, bytes] = {}
        transplanted = 0
        for sheet_name, source_part in source_parts.items():
            source_root = ElementTree.fromstring(source_zip.read(source_part))
            calculated_root = ElementTree.fromstring(
                calculated_zip.read(calculated_parts[sheet_name])
            )
            calculated_cells = {
                cell.attrib["r"]: cell
                for cell in calculated_root.findall(f".//{{{SHEET_NS}}}c")
                if cell.find(f"{{{SHEET_NS}}}f") is not None
            }
            for cell in source_root.findall(f".//{{{SHEET_NS}}}c"):
                if cell.find(f"{{{SHEET_NS}}}f") is None:
                    continue
                coordinate = cell.attrib["r"]
                calculated_cell = calculated_cells.get(coordinate)
                calculated_value = (
                    calculated_cell.find(f"{{{SHEET_NS}}}v")
                    if calculated_cell is not None
                    else None
                )
                if calculated_value is None:
                    raise RuntimeError(
                        f"LibreOffice did not cache formula {sheet_name}!{coordinate}"
                    )
                source_value = cell.find(f"{{{SHEET_NS}}}v")
                if source_value is None:
                    source_value = ElementTree.SubElement(cell, f"{{{SHEET_NS}}}v")
                source_value.text = calculated_value.text
                if calculated_cell is not None and "t" in calculated_cell.attrib:
                    cell.attrib["t"] = calculated_cell.attrib["t"]
                else:
                    cell.attrib.pop("t", None)
                transplanted += 1
            replacements[source_part] = ElementTree.tostring(
                source_root, encoding="utf-8", xml_declaration=True
            )

        with zipfile.ZipFile(output, "w") as output_zip:
            for item in source_zip.infolist():
                data = replacements.get(item.filename, source_zip.read(item.filename))
                output_zip.writestr(copy(item), data)
    return transplanted


def libreoffice_version(soffice: str) -> str:
    try:
        result = subprocess.run(
            [soffice, "--version"],
            capture_output=True,
            text=True,
            timeout=15,
        )
    except (OSError, subprocess.TimeoutExpired):
        return "unknown"
    return (result.stdout or result.stderr).strip().splitlines()[0][:200] or "unknown"


def recalculate(input_path: Path, output_path: Path, timeout: int) -> dict[str, Any]:
    if input_path.resolve() == output_path.resolve():
        raise ValueError("input and output paths must be different")
    if input_path.suffix.lower() != ".xlsx" or output_path.suffix.lower() != ".xlsx":
        raise ValueError("recalculation accepts .xlsx input and output paths only")
    if not input_path.is_file() or input_path.stat().st_size <= 0:
        raise ValueError("input workbook must be a non-empty regular file")
    check_zip(input_path)
    input_report = workbook_report(input_path)

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeUnavailableError("LibreOffice is missing from the baked runtime")

    with tempfile.TemporaryDirectory(prefix="xlsx-recalc-") as temporary:
        root = Path(temporary)
        input_dir = root / "input"
        output_dir = root / "output"
        profile_dir = root / "profile"
        input_dir.mkdir()
        output_dir.mkdir()
        profile_user_dir = profile_dir / "user"
        profile_user_dir.mkdir(parents=True)
        (profile_user_dir / "registrymodifications.xcu").write_text(
            PROFILE_XML, encoding="utf-8"
        )
        staged_input = input_dir / "workbook.xlsx"
        shutil.copyfile(input_path, staged_input)

        command = [
            soffice,
            f"-env:UserInstallation={profile_dir.resolve().as_uri()}",
            "--headless",
            "--nologo",
            "--nodefault",
            "--norestore",
            "--infilter=Calc MS Excel 2007 XML",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(output_dir),
            str(staged_input),
        ]
        try:
            process = subprocess.run(command, capture_output=True, text=True, timeout=timeout)
        except subprocess.TimeoutExpired as error:
            raise RuntimeError(f"LibreOffice recalculation timed out after {timeout}s") from error
        if process.returncode != 0:
            raise RuntimeError(
                f"LibreOffice exited {process.returncode}: "
                f"{(process.stderr or process.stdout).strip()[:1000]}"
            )

        produced = output_dir / "workbook.xlsx"
        check_zip(produced)
        output_report = workbook_report(produced)
        if output_report["formulas"] != input_report["formulas"]:
            input_cells = set(input_report["formulas"])
            output_cells = set(output_report["formulas"])
            missing = sorted(input_cells - output_cells)
            added = sorted(output_cells - input_cells)
            changed = sorted(
                cell
                for cell in input_cells & output_cells
                if input_report["formulas"][cell] != output_report["formulas"][cell]
            )
            raise RuntimeError(
                "LibreOffice changed formula lineage: "
                f"missing={missing[:5]} added={added[:5]} changed={changed[:5]}"
            )
        if output_report["formulaCells"] > 0 and output_report["missingCachedValues"]:
            sample = ", ".join(output_report["missingCachedValues"][:10])
            raise RuntimeError(f"recalculation left formula caches empty: {sample}")
        if output_report["cachedErrorValues"]:
            sample = ", ".join(output_report["cachedErrorValues"][:10])
            raise RuntimeError(f"recalculation produced cached formula errors: {sample}")
        input_sha256 = sha256(input_path)
        output_sha256 = sha256(produced)
        if output_report["formulaCells"] > 0 and input_sha256 == output_sha256:
            raise RuntimeError(
                "LibreOffice returned byte-identical formula-bearing input; "
                "recalculation was not proven"
            )

        output_path.parent.mkdir(parents=True, exist_ok=True)
        atomic_path = output_path.with_name(f".{output_path.name}.tmp")
        atomic_path.unlink(missing_ok=True)
        transplanted = transplant_formula_caches(input_path, produced, atomic_path)
        if transplanted != output_report["formulaCells"]:
            raise RuntimeError(
                "formula cache transplant count mismatch: "
                f"expected {output_report['formulaCells']}, received {transplanted}"
            )
        atomic_path.replace(output_path)

    check_zip(output_path)
    final_report = workbook_report(output_path)
    report = {
        "ok": True,
        "engine": "libreoffice",
        "engineVersion": libreoffice_version(soffice),
        "input": str(input_path),
        "output": str(output_path),
        "inputSha256": sha256(input_path),
        "outputSha256": sha256(output_path),
        "formulaCells": final_report["formulaCells"],
        "cachedValues": final_report["cachedValues"],
        "missingCachedValues": final_report["missingCachedValues"],
        "cachedErrorValues": final_report["cachedErrorValues"],
        "unsupportedFunctions": [],
    }
    report["bytesChanged"] = report["inputSha256"] != report["outputSha256"]
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--timeout", type=int, default=120)
    args = parser.parse_args()

    try:
        report = recalculate(args.input, args.output, max(1, args.timeout))
    except RuntimeUnavailableError as error:
        print(json.dumps({"ok": False, "kind": "runtime-unavailable", "error": str(error)}))
        return 2
    except Exception as error:
        print(json.dumps({"ok": False, "kind": "recalculation-failed", "error": str(error)}))
        return 1
    print(json.dumps(report, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
