#!/usr/bin/env python3
# SPDX-License-Identifier: MIT
"""Read-only workbook/CSV structure and formula-cache inspection."""

from __future__ import annotations

import argparse
import json
import sys
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any


def json_value(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def inspect_excel(path: Path, requested_sheet: str | None, sample_rows: int) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("openpyxl is missing from the baked runtime") from error

    keep_vba = path.suffix.lower() == ".xlsm"
    formulas = load_workbook(
        path,
        read_only=False,
        data_only=False,
        keep_links=True,
        keep_vba=keep_vba,
    )
    cached = load_workbook(
        path,
        read_only=True,
        data_only=True,
        keep_links=True,
        keep_vba=keep_vba,
    )
    selected = [requested_sheet] if requested_sheet else list(formulas.sheetnames)
    unknown = [name for name in selected if name not in formulas.sheetnames]
    if unknown:
        raise ValueError(f"unknown worksheet: {unknown[0]}")

    sheet_reports: list[dict[str, Any]] = []
    total_formulas = 0
    total_missing_caches = 0
    for name in selected:
        formula_sheet = formulas[name]
        cached_sheet = cached[name]
        if not hasattr(formula_sheet, "iter_rows"):
            sheet_reports.append(
                {
                    "name": name,
                    "type": "chartsheet",
                    "state": formula_sheet.sheet_state,
                    "rows": 0,
                    "columns": 0,
                    "formulaCells": 0,
                    "missingCachedValues": [],
                    "tables": 0,
                    "charts": 1,
                    "mergedRanges": [],
                    "freezePanes": None,
                    "sampleRows": [],
                }
            )
            continue
        formula_count = 0
        missing_caches: list[str] = []
        samples: list[list[Any]] = []
        for row_index, (formula_row, cached_row) in enumerate(
            zip(formula_sheet.iter_rows(), cached_sheet.iter_rows()),
            start=1,
        ):
            if row_index <= sample_rows:
                samples.append([json_value(cell.value) for cell in formula_row[:20]])
            for formula_cell, cached_cell in zip(formula_row, cached_row):
                if formula_cell.data_type != "f":
                    continue
                formula_count += 1
                if cached_cell.value is None:
                    missing_caches.append(formula_cell.coordinate)
        total_formulas += formula_count
        total_missing_caches += len(missing_caches)
        sheet_reports.append(
            {
                "name": name,
                "type": "worksheet",
                "state": formula_sheet.sheet_state,
                "rows": formula_sheet.max_row,
                "columns": formula_sheet.max_column,
                "formulaCells": formula_count,
                "missingCachedValues": missing_caches[:100],
                "tables": len(formula_sheet.tables),
                "charts": len(getattr(formula_sheet, "_charts", [])),
                "mergedRanges": [
                    str(item) for item in list(formula_sheet.merged_cells.ranges)[:100]
                ],
                "freezePanes": str(formula_sheet.freeze_panes or "") or None,
                "sampleRows": samples,
            }
        )

    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
        return {
            "format": path.suffix.lower().lstrip("."),
            "path": str(path),
            "sheets": sheet_reports,
            "summary": {
                "sheetCount": len(formulas.sheetnames),
                "formulaCells": total_formulas,
                "missingCachedValues": total_missing_caches,
                "chartParts": sum(
                    name.startswith("xl/charts/") and name.endswith(".xml")
                    for name in names
                ),
                "tableParts": sum(
                    name.startswith("xl/tables/") and name.endswith(".xml")
                    for name in names
                ),
                "hasVba": "xl/vbaProject.bin" in names,
                "hasExternalLinks": any(
                    name.startswith("xl/externalLinks/") for name in names
                ),
            },
        }
    finally:
        formulas.close()
        cached.close()


def inspect_delimited(path: Path, sample_rows: int) -> dict[str, Any]:
    try:
        import pandas as pd
    except ImportError as error:
        raise RuntimeError("pandas is missing from the baked runtime") from error

    separator = "\t" if path.suffix.lower() == ".tsv" else ","
    errors: list[str] = []
    frame = None
    encoding = None
    for candidate in ("utf-8-sig", "utf-8", "gbk", "latin-1"):
        try:
            frame = pd.read_csv(path, sep=separator, encoding=candidate)
            encoding = candidate
            break
        except (UnicodeDecodeError, pd.errors.ParserError, ValueError) as error:
            errors.append(f"{candidate}: {error}")
    if frame is None:
        raise ValueError("unable to read delimited file; " + " | ".join(errors))
    return {
        "format": path.suffix.lower().lstrip("."),
        "path": str(path),
        "encoding": encoding,
        "rows": int(len(frame.index)),
        "columns": [str(column) for column in frame.columns],
        "nulls": {str(key): int(value) for key, value in frame.isna().sum().items()},
        "duplicateRows": int(frame.duplicated().sum()),
        "sampleRows": [
            {str(key): json_value(value) for key, value in row.items()}
            for row in frame.head(sample_rows).to_dict(orient="records")
        ],
    }


def render_human(report: dict[str, Any]) -> str:
    lines = [f"Format: {report['format']}", f"Path: {report['path']}"]
    if "sheets" in report:
        summary = report["summary"]
        lines.append(
            f"Sheets: {summary['sheetCount']} | formulas: {summary['formulaCells']} | "
            f"missing caches: {summary['missingCachedValues']} | charts: {summary['chartParts']}"
        )
        for sheet in report["sheets"]:
            lines.append(
                f"- {sheet['name']}: {sheet['rows']} rows x {sheet['columns']} columns; "
                f"formulas={sheet['formulaCells']} missingCaches={len(sheet['missingCachedValues'])}"
            )
    else:
        lines.append(
            f"Rows: {report['rows']} | columns: {len(report['columns'])} | "
            f"duplicates: {report['duplicateRows']} | encoding: {report['encoding']}"
        )
    return "\n".join(lines)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--sheet")
    parser.add_argument("--sample-rows", type=int, default=5)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    try:
        if not args.input.is_file() or args.input.stat().st_size <= 0:
            raise ValueError("input must be a non-empty regular file")
        suffix = args.input.suffix.lower()
        if suffix in {".xlsx", ".xlsm"}:
            report = inspect_excel(args.input, args.sheet, max(0, args.sample_rows))
        elif suffix in {".csv", ".tsv"}:
            if args.sheet:
                raise ValueError("--sheet applies only to XLSX/XLSM inputs")
            report = inspect_delimited(args.input, max(0, args.sample_rows))
        else:
            raise ValueError("supported inputs are .xlsx, .xlsm, .csv, and .tsv")
    except Exception as error:
        print(json.dumps({"ok": False, "error": str(error)}), file=sys.stderr)
        return 1

    print(json.dumps(report, indent=2) if args.json else render_human(report))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
